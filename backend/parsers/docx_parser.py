"""DOCX parser (python-docx): paragraphs and table rows, all with locations."""
from __future__ import annotations

from pathlib import Path

import docx


def parse(path: str | Path) -> dict:
    document = docx.Document(str(path))
    blocks: list[dict] = []

    for i, para in enumerate(document.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        blocks.append({"location": f"para_{i}", "text": text, "meta": {}})

    for t, table in enumerate(document.tables):
        for r, row in enumerate(table.rows):
            cells = [c.text.strip() for c in row.cells]
            if not any(cells):
                continue
            blocks.append(
                {
                    "location": f"tbl_{t}_row_{r}",
                    "text": " | ".join(cells),
                    "meta": {"cells": cells},
                }
            )

    return {"artifact_type": "docx", "blocks": blocks}
