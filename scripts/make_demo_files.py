"""Generate the binary demo files (PRD.docx / test_plan.docx / release_plan.xlsx).

Kept as a script (not committed binaries) so the dataset stays reviewable.
Run:  .venv/bin/python scripts/make_demo_files.py
"""
from __future__ import annotations

from pathlib import Path

import docx
import openpyxl

DEMO_DIR = Path(__file__).resolve().parent.parent / "demo" / "workspace_alpha"


def make_prd() -> None:
    document = docx.Document()
    document.add_heading("Alpha V2.0 产品需求文档", level=0)
    document.add_paragraph("1. 概述")
    document.add_paragraph("Alpha V2.0 将于 2026 年 9 月 20 日正式发布。")
    document.add_paragraph("2. 范围")
    document.add_paragraph("- 支付模块接入新接口。")
    document.add_paragraph("- 个人中心改版。")
    document.add_paragraph("3. 里程碑")
    document.add_paragraph("- 负责人：张伟（整体交付）。")
    document.add_paragraph("- 本需求不含负责人与状态字段的事实同步（MVP 范围外）。")
    document.save(DEMO_DIR / "PRD.docx")


def make_test_plan() -> None:
    document = docx.Document()
    document.add_heading("Alpha V2.0 测试计划", level=0)
    document.add_paragraph("回归测试将于 9 月 18 日完成，随后输出测试报告。")
    document.add_paragraph("测试范围：支付模块、个人中心改版。")
    document.save(DEMO_DIR / "test_plan.docx")


def make_release_plan() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Release Plan"
    rows = [
        ("Project", "Alpha V2.0"),
        ("Release Date", "2026-09-20"),
        ("Regression Deadline", "2026-09-18"),
        ("Status", "Testing"),
    ]
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 16
    workbook.save(DEMO_DIR / "release_plan.xlsx")


if __name__ == "__main__":
    DEMO_DIR.mkdir(parents=True, exist_ok=True)
    make_prd()
    make_test_plan()
    make_release_plan()
    print("demo files written to", DEMO_DIR)
