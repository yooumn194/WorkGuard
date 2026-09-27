"""Feishu integration tests — fully mocked: no network access ever."""
import json

import httpx
import pytest
from fastapi.testclient import TestClient

from backend.config import settings
from backend.db import Base, SessionLocal, engine, init_db
from backend.integrations.feishu import service as feishu_service
from backend.integrations.feishu.client import FeishuClient, FeishuError
from backend.main import app
from backend.services.ingest import create_workspace


class FakeFeishuClient(FeishuClient):
    """Canned drive/bitable responses; asserts no real HTTP happens."""

    def __init__(self, files=None, records=None, doc_bytes=b"dummy-docx", table_id="tbl_1"):
        super().__init__(app_id="fake", app_secret="fake")
        self.files = files or []
        self.records = records or []
        self.doc_bytes = doc_bytes
        self.table_id = table_id
        self.export_called = 0
        self.subscriptions = []

    def list_folder_files(self, folder_token: str) -> list[dict]:
        return self.files

    def export_doc_as_docx(self, file_token: str, timeout_seconds: int = 20) -> bytes:
        self.export_called += 1
        return self.doc_bytes

    def list_bitable_records(self, app_token: str, table_id: str) -> list[dict]:
        return self.records

    def list_bitable_tables(self, app_token: str) -> list[dict]:
        return [{"table_id": self.table_id}]

    def subscribe_file_events(self, file_token: str, file_type: str) -> dict:
        self.subscriptions.append((file_token, file_type))
        return {"code": 0, "msg": "success", "data": {}}


@pytest.fixture(autouse=True)
def clean_db():
    Base.metadata.drop_all(engine)
    init_db()
    yield


def test_bitable_records_flatten_to_markdown():
    records = [
        {"record_id": "r1", "fields": {"Project": "Alpha V2.0", "Release Date": "2026-09-20"}},
        {"record_id": "r2", "fields": {"Project": "Alpha V2.0", "Release Date": {"text": "2026-09-27"},
                                        "Owner": [{"name": "张伟"}]}},
    ]
    markdown = FeishuClient.bitable_records_to_markdown("Release Plan", records)
    assert "Release Date: 2026-09-20" in markdown
    assert "Release Date: 2026-09-27" in markdown
    assert "Owner: 张伟" in markdown


def test_sync_workspace_pulls_docs_and_bitables():
    import io

    import docx as docx_lib

    buffer = io.BytesIO()
    document = docx_lib.Document()
    document.add_paragraph("Alpha V2.0 将于 2026 年 9 月 20 日正式发布。")
    document.save(buffer)
    fake = FakeFeishuClient(
        doc_bytes=buffer.getvalue(),
        files=[
            {"token": "doc_1", "name": "PRD", "type": "docx", "url": ""},
            {"token": "bit_1", "name": "Release Plan", "type": "bitable", "url": ""},
            {"token": "img_1", "name": "logo", "type": "file", "url": ""},
        ],
        records=[{"record_id": "r1", "fields": {"Project": "Alpha V2.0",
                                                 "Release Date": "2026-09-20"}}],
    )
    monkey_table = lambda client, app_token: "tbl_1"  # noqa: E731
    original = feishu_service._first_table_id
    feishu_service._first_table_id = monkey_table
    try:
        with SessionLocal() as session:
            workspace = create_workspace(
                session, "feishu-ws",
                preset_entities=[{"canonical_name": "Alpha V2.0", "aliases": ["Alpha"]}],
            )
            result = feishu_service.sync_workspace(session, workspace.id, "fld_x", client=fake)
            repeated = feishu_service.sync_workspace(session, workspace.id, "fld_x", client=fake)
    finally:
        feishu_service._first_table_id = original

    assert [item["name"] for item in result["synced"]] == ["PRD.docx", "Release Plan.md"]
    assert result["skipped"] == [{"name": "logo", "type": "file", "reason": "unsupported type"}]
    assert len(result["runs"]) == 2  # change detection ran per artifact
    assert all(run["run_state"] is not None for run in result["runs"])
    assert result["event_subscriptions"] == [
        {"artifact_id": result["synced"][0]["artifact_id"], "name": "PRD.docx",
         "subscribed": True, "error": ""},
        {"artifact_id": result["synced"][1]["artifact_id"], "name": "Release Plan.md",
         "subscribed": True, "error": ""},
    ]
    assert fake.subscriptions[:2] == [("doc_1", "docx"), ("bit_1", "bitable")]
    assert repeated["synced"] == []
    assert len(repeated["unchanged"]) == 2
    assert repeated["runs"] == []

    document = docx_lib.Document()
    changed_buffer = io.BytesIO()
    document.add_paragraph("Alpha V2.0 将于 2026 年 9 月 27 日正式发布。")
    document.save(changed_buffer)
    fake.doc_bytes = changed_buffer.getvalue()
    with SessionLocal() as session:
        changed = feishu_service.sync_workspace(session, workspace.id, "fld_x", client=fake)
        from sqlalchemy import select

        from backend.models import Artifact
        artifacts = session.scalars(select(Artifact).where(Artifact.workspace_id == workspace.id)).all()
    assert len(changed["synced"]) == 1
    assert len(changed["unchanged"]) == 1
    assert len(artifacts) == 2
    doc_artifact = next(item for item in artifacts if item.name == "PRD.docx")
    assert doc_artifact.current_version == 2


def test_sync_requires_configuration():
    with SessionLocal() as session:
        workspace = create_workspace(session, "feishu-ws-2")
        unconfigured = FeishuClient(app_id="", app_secret="")
        with pytest.raises(FeishuError):
            feishu_service.sync_workspace(session, workspace.id, "fld_x", client=unconfigured)


def test_sync_exposes_event_subscription_failure_without_blocking_manual_sync():
    import io

    import docx as docx_lib

    buffer = io.BytesIO()
    document = docx_lib.Document()
    document.add_paragraph("Alpha V2.0 将于 2026 年 9 月 20 日正式发布。")
    document.save(buffer)
    fake = FakeFeishuClient(
        doc_bytes=buffer.getvalue(),
        files=[{"token": "doc_1", "name": "PRD", "type": "docx", "url": ""}],
    )

    def fail_subscription(file_token: str, file_type: str) -> dict:
        raise FeishuError("missing docs:event:subscribe")

    fake.subscribe_file_events = fail_subscription  # type: ignore[method-assign]
    with SessionLocal() as session:
        workspace = create_workspace(session, "manual-sync-without-events")
        result = feishu_service.sync_workspace(session, workspace.id, "fld_x", client=fake)

    assert len(result["synced"]) == 1
    assert result["event_subscriptions"] == [{
        "artifact_id": result["synced"][0]["artifact_id"],
        "name": "PRD.docx",
        "subscribed": False,
        "error": "missing docs:event:subscribe",
    }]


def test_webhook_challenge_and_status_endpoints(monkeypatch):
    monkeypatch.setattr(settings, "feishu_verification_token", "verify-me")
    monkeypatch.setattr(settings, "feishu_app_id", "")
    monkeypatch.setattr(settings, "feishu_app_secret", "")
    with SessionLocal() as session:
        workspace = create_workspace(session, "unconfigured-feishu")
        workspace_id = workspace.id
    with TestClient(app) as client:
        response = client.post("/api/integrations/feishu/webhook",
                               json={"challenge": "ajls384kdd", "type": "url_verification",
                                     "token": "verify-me"})
        assert response.status_code == 200
        assert response.json() == {"challenge": "ajls384kdd"}

        status = client.get("/api/integrations/feishu/status")
        assert status.status_code == 200
        body = status.json()
        assert body["configured"] is False

        # sync endpoint refuses cleanly when unconfigured
        result = client.post("/api/integrations/feishu/sync",
                             json={"workspace_id": workspace_id, "folder_token": "fld"})
        assert result.status_code == 503
        assert "not configured" in result.json()["detail"]


def test_webhook_ignores_unrelated_events(monkeypatch):
    monkeypatch.setattr(settings, "feishu_verification_token", "verify-me")
    with TestClient(app) as client:
        payload = {"header": {"event_type": "im.message.receive_v1", "event_id": "evt-ignore",
                              "token": "verify-me"}, "event": {}}
        response = client.post("/api/integrations/feishu/webhook", json=payload)
        assert response.status_code == 200
        body = response.json()
        assert body["handled"] is False


def test_webhook_rejects_missing_or_wrong_verification_token(monkeypatch):
    monkeypatch.setattr(settings, "feishu_verification_token", "verify-me")
    with TestClient(app) as client:
        for token in (None, "wrong"):
            payload = {"challenge": "x"}
            if token:
                payload["token"] = token
            response = client.post("/api/integrations/feishu/webhook", json=payload)
            assert response.status_code == 403


def test_event_receipt_is_idempotent_and_routes_by_remote_file(monkeypatch):
    from backend.integrations.feishu.service import accept_webhook_event, create_binding
    from backend.models import BackgroundJob, FeishuRemoteFile, uid

    monkeypatch.setattr(settings, "feishu_verification_token", "verify-me")
    with SessionLocal() as session:
        workspace = create_workspace(session, "event-routing")
        binding = create_binding(session, workspace.id, "fld_1")
        session.add(FeishuRemoteFile(
            id=uid("fsf"), binding_id=binding.id, remote_token="doc_1",
            remote_type="docx", remote_name="PRD", artifact_id="art_1",
        ))
        session.commit()
        payload = {
            "header": {"event_type": "drive.file.edit_v1", "event_id": "evt-1",
                       "token": "verify-me"},
            "event": {"file_token": "doc_1"},
        }
        first = accept_webhook_event(payload, session)
        second = accept_webhook_event(payload, session)
        assert first["handled"] is True
        assert first["event_record_id"]
        job = session.get(BackgroundJob, first["job_id"])
        assert job is not None
        assert job.kind == "feishu_event"
        assert job.workspace_id == workspace.id
        assert job.payload == {"event_record_id": first["event_record_id"]}
        assert second == {"handled": False, "duplicate": True, "event_id": "evt-1"}


def test_approval_notification_payload_shape(monkeypatch):
    from sqlalchemy import select

    from backend.models import AuditLog

    sent = []

    def fake_send(self, text: str) -> dict:
        sent.append(text)
        return {"code": 0}

    monkeypatch.setattr(FeishuClient, "send_webhook_text", fake_send)
    monkeypatch.setattr(FeishuClient, "webhook_configured", property(lambda self: True))
    monkeypatch.setattr(settings, "workguard_public_base_url", "https://workguard.test")
    with SessionLocal() as session:
        workspace = create_workspace(session, "notification-audit")
    card = {
        "workspace_id": workspace.id, "change_id": "chg-test",
        "entity": "Alpha V2.0", "predicate": "release_date",
        "old_value": "2026-09-20", "new_value": "2026-09-27",
        "source": {"artifact": "weekly_0905.md", "evidence": "由 9 月 20 日调整至 9 月 27 日"},
        "conflicts": [{"verdict": "conflict"}, {"verdict": "conflict"}],
        "impacts": [{"id": "i1"}],
    }
    assert feishu_service.notify_pending_approval(card) is True
    assert len(sent) == 1
    assert "2026-09-20 -> 2026-09-27" in sent[0]
    assert "2" in sent[0]  # conflict count included
    assert (
        "https://workguard.test/app/?workspace="
        f"{workspace.id}&change=chg-test"
    ) in sent[0]
    with SessionLocal() as session:
        audit = session.scalars(
            select(AuditLog).where(AuditLog.tool == "feishu_approval_notification")
        ).one()
        assert audit.status == "success"
        assert audit.output == {"sent": True, "error": ""}


def test_http_client_contract_and_webhook_error_code():
    requests = []

    def handler(request: httpx.Request):
        requests.append(request)
        if request.url.path.endswith("/tenant_access_token/internal"):
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "tenant-x", "expire": 7200})
        if request.url.path.endswith("/drive/v1/files"):
            return httpx.Response(200, json={"code": 0, "data": {"files": [], "has_more": False}})
        return httpx.Response(200, json={"code": 19001, "msg": "rejected"})

    transport = httpx.MockTransport(handler)
    http = httpx.Client(transport=transport)
    client = FeishuClient(app_id="app-x", app_secret="secret-x", base_url="https://unit.test",
                          webhook_url="https://bot.test/hook", http_client=http)
    assert client.tenant_access_token() == "tenant-x"
    assert client.list_folder_files("fld-x") == []
    assert requests[1].headers["authorization"] == "Bearer tenant-x"
    with pytest.raises(FeishuError, match="rejected"):
        client.send_webhook_text("hello")


def test_root_drive_sentinel_omits_folder_token():
    requests = []

    def handler(request: httpx.Request):
        requests.append(request)
        if request.url.path.endswith("/tenant_access_token/internal"):
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "tenant-x"})
        return httpx.Response(200, json={"code": 0, "data": {"files": [], "has_more": False}})

    client = FeishuClient(
        app_id="app-x", app_secret="secret-x", base_url="https://unit.test",
        http_client=httpx.Client(transport=httpx.MockTransport(handler)),
    )
    assert client.list_folder_files("root") == []
    assert "folder_token" not in requests[1].url.params


def test_export_uses_real_file_token_download_contract():
    requests = []

    def handler(request: httpx.Request):
        requests.append(request)
        path = request.url.path
        if path.endswith("/tenant_access_token/internal"):
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "tenant-x"})
        if path.endswith("/export_tasks"):
            return httpx.Response(200, json={"code": 0, "data": {"ticket": "ticket-x"}})
        if path.endswith("/export_tasks/ticket-x"):
            return httpx.Response(200, json={
                "code": 0, "data": {"result": {"job_status": 0, "file_token": "file-x"}},
            })
        if path.endswith("/export_tasks/file/file-x/download"):
            return httpx.Response(200, content=b"docx-bytes")
        return httpx.Response(404)

    client = FeishuClient(
        app_id="app-x", app_secret="secret-x", base_url="https://unit.test",
        http_client=httpx.Client(transport=httpx.MockTransport(handler)),
    )
    assert client.export_doc_as_docx("doc-x") == b"docx-bytes"
    assert requests[-1].headers["authorization"] == "Bearer tenant-x"


def test_subscribe_file_events_uses_per_document_contract():
    requests = []

    def handler(request: httpx.Request):
        requests.append(request)
        if request.url.path.endswith("/tenant_access_token/internal"):
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "tenant-x"})
        return httpx.Response(200, json={"code": 0, "msg": "success", "data": {}})

    client = FeishuClient(
        app_id="app-x", app_secret="secret-x", base_url="https://unit.test",
        http_client=httpx.Client(transport=httpx.MockTransport(handler)),
    )
    assert client.subscribe_file_events("doc-x", "docx")["code"] == 0
    request = requests[-1]
    assert request.method == "POST"
    assert request.url.path.endswith("/drive/v1/files/doc-x/subscribe")
    assert request.url.params["file_type"] == "docx"
    assert request.headers["authorization"] == "Bearer tenant-x"


def test_app_bot_notification_contract():
    requests = []

    def handler(request: httpx.Request):
        requests.append(request)
        if request.url.path.endswith("/tenant_access_token/internal"):
            return httpx.Response(200, json={"code": 0, "tenant_access_token": "tenant-x"})
        return httpx.Response(200, json={"code": 0, "msg": "success", "data": {}})

    client = FeishuClient(
        app_id="app-x", app_secret="secret-x", base_url="https://unit.test",
        notify_receive_id="chat-x", notify_receive_id_type="chat_id",
        http_client=httpx.Client(transport=httpx.MockTransport(handler)),
    )
    assert client.send_app_text("approval pending")["code"] == 0
    request = requests[-1]
    assert request.url.params["receive_id_type"] == "chat_id"
    payload = json.loads(request.content)
    assert payload["receive_id"] == "chat-x"
    assert json.loads(payload["content"])["text"] == "approval pending"


def test_sync_orders_baseline_before_meeting_decision():
    from backend.integrations.feishu.service import _sync_priority

    items = [
        {"name": "WorkGuard E2E - 周会改期决议"},
        {"name": "WorkGuard E2E - PRD 基线"},
        {"name": "Release Plan"},
    ]
    assert [item["name"] for item in sorted(items, key=_sync_priority)] == [
        "WorkGuard E2E - PRD 基线",
        "Release Plan",
        "WorkGuard E2E - 周会改期决议",
    ]
