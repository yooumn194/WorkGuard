"""Write-back tool tests: markdown direct write, docx/xlsx controlled write,
suggestion patches, post verification and rollback."""
from pathlib import Path

import docx
import openpyxl
import pytest
from sqlalchemy import select

from backend.db import SessionLocal, init_db
from backend.models import Artifact, ArtifactVersion, Workspace
from backend.services.ingest import upload_artifact
from backend.tools import document_tools as dt


def _ensure_workspace(session, workspace_id: str) -> None:
    if session.get(Workspace, workspace_id) is None:
        session.add(Workspace(id=workspace_id, name=workspace_id))
        session.flush()


@pytest.fixture()
def artifacts(tmp_path):
    init_db()

    md = tmp_path / "launch_plan.md"
    md.write_text("# Alpha V2.0 上线计划\n\nAlpha V2.0 定于 9 月 20 日正式上线，灰度发布从 9 月 18 日开始。\n")

    document = docx.Document()
    document.add_paragraph("Alpha V2.0 将于 2026 年 9 月 20 日正式发布。")
    docx_path = tmp_path / "PRD.docx"
    document.save(docx_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(("Release Date", "2026-09-20"))
    ws.append(("Formula", '=B1&"!"'))
    xlsx_path = tmp_path / "plan.xlsx"
    wb.save(xlsx_path)

    with SessionLocal() as session:
        _ensure_workspace(session, "ws_w")
        md_art = upload_artifact(session, "ws_w", "launch_plan.md", md.read_bytes())
        docx_art = upload_artifact(session, "ws_w", "PRD.docx", docx_path.read_bytes())
        xlsx_art = upload_artifact(session, "ws_w", "plan.xlsx", xlsx_path.read_bytes())
        session.commit()
        yield {"md": md_art, "docx": docx_art, "xlsx": xlsx_art}


def _previous_version(session, artifact, after_version_number):
    return session.scalars(
        select(ArtifactVersion)
        .where(
            ArtifactVersion.artifact_id == artifact.id,
            ArtifactVersion.version < after_version_number,
        )
        .order_by(ArtifactVersion.version.desc())
    ).first()


def test_markdown_direct_write_and_rollback(artifacts):
    md = artifacts["md"]
    edits = [{"location": "line_3", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}]

    outcome = dt.write_markdown(md, edits)
    assert outcome["success"]
    assert outcome["result"]["applied"][0]["new"] == "9 月 27 日"
    text = Path(md.source_path).read_text()
    assert "9 月 27 日" in text and "9 月 20 日" not in text

    report = dt.post_verify(md, "2026-09-20", "2026-09-27", ["line_3"])
    assert report["success"]

    # version after write (v2), rollback to v1 restores the original text
    with SessionLocal() as session:
        version = dt.commit_version(session, md)
        previous = _previous_version(session, md, version.version)
        restored = dt.restore_version(session, md, previous)
        session.commit()
    assert restored["success"]
    assert "9 月 20 日" in Path(md.source_path).read_text()


def test_docx_controlled_write_single_run(artifacts):
    docx_art = artifacts["docx"]
    edits = [{"location": "", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}]
    outcome = dt.write_docx(docx_art, edits)
    assert outcome["success"], outcome
    assert outcome["result"]["applied"][0]["new"] == "2026 年 9 月 27 日"
    document = docx.Document(docx_art.source_path)
    assert any("2026 年 9 月 27 日" in p.text for p in document.paragraphs)


def test_docx_controlled_write_inside_table_row(tmp_path):
    init_db()
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Release Date"
    table.cell(0, 1).text = "2026-09-20"
    path = tmp_path / "table.docx"
    document.save(path)
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_docx_table")
        artifact = upload_artifact(session, "ws_docx_table", path.name, path.read_bytes())
        session.commit()
    outcome = dt.write_docx(
        artifact,
        [{"location": "tbl_0_row_0", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}],
    )
    assert outcome["success"] is True
    reloaded = docx.Document(artifact.source_path)
    assert reloaded.tables[0].cell(0, 1).text == "2026-09-27"


def test_docx_report_needs_manual_when_spans_runs(tmp_path):
    init_db()
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("上线日期 2026-09-")
    paragraph.add_run("20")  # date deliberately split across two runs
    path = tmp_path / "split.docx"
    document.save(path)
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_w2")
        artifact = upload_artifact(session, "ws_w2", "split.docx", path.read_bytes())
        session.commit()

    outcome = dt.write_docx(artifact, [{"location": "", "old_iso": "2026-09-20",
                                        "new_iso": "2026-09-27"}])
    # controlled mode refuses: file must be untouched
    reloaded = docx.Document(artifact.source_path)
    assert reloaded.paragraphs[0].text == "上线日期 2026-09-20"
    assert outcome["success"] is False or outcome["result"]["refused"]


def test_xlsx_write_and_formula_protection(artifacts):
    xlsx = artifacts["xlsx"]
    edits = [{"location": "Sheet!row_1", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}]
    outcome = dt.write_xlsx(xlsx, edits)
    assert outcome["success"]
    wb = openpyxl.load_workbook(xlsx.source_path)
    assert wb.active["B1"].value == "2026-09-27"
    assert wb.active["B2"].value.startswith("=")  # formula cell untouched


def test_xlsx_date_cell_hidden_sheet_merged_cell_and_formula_recalculation(tmp_path):
    import datetime as datetime_lib

    init_db()
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Cover"
    schedule = wb.create_sheet("Hidden Schedule")
    schedule.sheet_state = "hidden"
    schedule.merge_cells("A1:B1")
    schedule["A1"] = datetime_lib.date(2026, 9, 20)
    schedule["A1"].number_format = "yyyy-mm-dd"
    schedule["C1"] = "=A1+7"
    path = tmp_path / "hidden.xlsx"
    wb.save(path)
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_xlsx_hidden")
        artifact = upload_artifact(session, "ws_xlsx_hidden", path.name, path.read_bytes())
        session.commit()

    outcome = dt.write_xlsx(
        artifact,
        [{"location": "Hidden Schedule!row_1", "old_iso": "2026-09-20",
          "new_iso": "2026-09-27"}],
    )
    assert outcome["success"] is True
    assert outcome["result"]["formula_recalculation"] == "on_open"
    reloaded = openpyxl.load_workbook(artifact.source_path, data_only=False)
    sheet = reloaded["Hidden Schedule"]
    assert sheet.sheet_state == "hidden"
    assert sheet["A1"].value.date().isoformat() == "2026-09-27"
    assert sheet["C1"].value == "=A1+7"
    assert reloaded.calculation.calcMode == "auto"
    assert reloaded.calculation.forceFullCalc is True


def test_xlsx_missing_sheet_or_row_is_refused_atomically(artifacts):
    xlsx = artifacts["xlsx"]
    before = Path(xlsx.source_path).read_bytes()
    missing_sheet = dt.write_xlsx(
        xlsx,
        [{"location": "Missing!row_1", "old_iso": "2026-09-20",
          "new_iso": "2026-09-27"}],
    )
    missing_row = dt.write_xlsx(
        xlsx,
        [{"location": "Sheet!row_999", "old_iso": "2026-09-20",
          "new_iso": "2026-09-27"}],
    )
    assert missing_sheet["success"] is False
    assert missing_row["success"] is False
    assert Path(xlsx.source_path).read_bytes() == before


def test_suggestion_patch_leaves_file_untouched(artifacts):
    docx_art = artifacts["docx"]
    original = Path(docx_art.source_path).read_bytes()
    edits = [{"location": "para_0", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}]
    blocks = dt.parse_file(docx_art.source_path, "docx")["blocks"]
    outcome = dt.generate_patch(docx_art, edits, blocks, reason="test")
    assert outcome["success"]
    assert Path(docx_art.source_path).read_bytes() == original  # unchanged!
    patch = Path(outcome["result"]["patch_path"]).read_text()
    assert "修改建议" in patch and "2026 年 9 月 27 日" in patch


def test_post_verify_detects_residual_old_value(artifacts):
    md = artifacts["md"]
    # simulate a partial write: new value added, old value still present
    path = Path(md.source_path)
    path.write_text("上线 9 月 20 日，另有口径 9 月 27 日。\n")
    report = dt.post_verify(md, "2026-09-20", "2026-09-27", [])
    assert report["success"] is False


def test_markdown_refuses_ambiguous_duplicate_and_is_atomic(tmp_path):
    init_db()
    source = tmp_path / "duplicate.md"
    original = "上线日期 9月20日；历史基线也是 9月20日。\n".encode()
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_duplicate")
        artifact = upload_artifact(session, "ws_duplicate", source.name, original)
        session.commit()
    before = Path(artifact.source_path).read_bytes()
    outcome = dt.write_markdown(
        artifact,
        [{"location": "line_1", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}],
    )
    assert outcome["success"] is False
    assert Path(artifact.source_path).read_bytes() == before


def test_markdown_does_not_match_short_date_inside_a_different_explicit_year(tmp_path):
    init_db()
    original = "历史发布日期：2025-09-20。\n".encode()
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_year_guard")
        artifact = upload_artifact(session, "ws_year_guard", "history.md", original)
        session.commit()
    before = Path(artifact.source_path).read_bytes()
    outcome = dt.write_markdown(
        artifact,
        [{"location": "line_1", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}],
    )
    assert outcome["success"] is False
    assert Path(artifact.source_path).read_bytes() == before


def test_markdown_preserves_bom_crlf_and_final_newline(tmp_path):
    init_db()
    original = b"\xef\xbb\xbf# Plan\r\nRelease: 2026-09-20\r\n"
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_encoding")
        artifact = upload_artifact(session, "ws_encoding", "encoded.md", original)
        session.commit()
    outcome = dt.write_markdown(
        artifact,
        [{"location": "line_2", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}],
    )
    updated = Path(artifact.source_path).read_bytes()
    assert outcome["success"] is True
    assert updated.startswith(b"\xef\xbb\xbf")
    assert b"\r\n" in updated and updated.endswith(b"\r\n")
    assert b"2026-09-27" in updated

    with SessionLocal() as session:
        persistent = session.get(Artifact, artifact.id)
        written = dt.commit_version(session, persistent)
        previous = _previous_version(session, persistent, written.version)
        assert dt.restore_version(session, persistent, previous)["success"]
        session.commit()
    assert Path(artifact.source_path).read_bytes() == original


def test_year_only_change_is_refused_when_short_surface_cannot_show_it(tmp_path):
    init_db()
    original = "上线日期：9月20日。\n".encode()
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_year_surface")
        artifact = upload_artifact(session, "ws_year_surface", "short.md", original)
        session.commit()
    outcome = dt.write_markdown(
        artifact,
        [{"location": "line_1", "old_iso": "2026-09-20", "new_iso": "2027-09-20"}],
    )
    assert outcome["success"] is False
    assert Path(artifact.source_path).read_bytes() == original


def test_non_utf8_text_is_refused_without_transcoding(tmp_path):
    init_db()
    original = "café Release Date 2026-09-20\n".encode("latin-1")
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_latin1")
        artifact = upload_artifact(session, "ws_latin1", "legacy.txt", original)
        session.commit()
    outcome = dt.write_markdown(
        artifact,
        [{"location": "line_1", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}],
    )
    assert outcome["success"] is False
    assert "non-UTF-8" in outcome["error"]
    assert Path(artifact.source_path).read_bytes() == original


def test_docx_multi_edit_is_atomic_when_one_target_is_ambiguous(tmp_path):
    init_db()
    document = docx.Document()
    document.add_paragraph("发布日期 2026-09-20")
    document.add_paragraph("复核 2026-09-20，同时保留 2026-09-20")
    path = tmp_path / "atomic.docx"
    document.save(path)
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_docx_atomic")
        artifact = upload_artifact(session, "ws_docx_atomic", path.name, path.read_bytes())
        session.commit()
    before = Path(artifact.source_path).read_bytes()
    outcome = dt.write_docx(artifact, [
        {"location": "para_0", "old_iso": "2026-09-20", "new_iso": "2026-09-27"},
        {"location": "para_1", "old_iso": "2026-09-20", "new_iso": "2026-09-27"},
    ])
    assert outcome["success"] is False
    assert Path(artifact.source_path).read_bytes() == before


def test_xlsx_refuses_multiple_matching_cells_in_target_row(tmp_path):
    init_db()
    wb = openpyxl.Workbook()
    wb.active.append(("当前", "2026-09-20", "历史", "2026-09-20"))
    source = tmp_path / "ambiguous.xlsx"
    wb.save(source)
    with SessionLocal() as session:
        _ensure_workspace(session, "ws_xlsx_ambiguous")
        artifact = upload_artifact(session, "ws_xlsx_ambiguous", source.name, source.read_bytes())
        session.commit()
    before = Path(artifact.source_path).read_bytes()
    outcome = dt.write_xlsx(
        artifact,
        [{"location": "Sheet!row_1", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}],
    )
    assert outcome["success"] is False
    assert Path(artifact.source_path).read_bytes() == before


def test_post_verify_does_not_accept_new_date_outside_target(artifacts):
    md = artifacts["md"]
    Path(md.source_path).write_text("目标仍是 9月20日。\n其他行已经是 9月27日。\n")
    report = dt.post_verify(md, "2026-09-20", "2026-09-27", ["line_1"])
    assert report["success"] is False


def test_restore_appends_monotonic_version_instead_of_reusing_number(artifacts):
    md = artifacts["md"]
    assert dt.write_markdown(md, [{"location": "line_3", "old_iso": "2026-09-20",
                                   "new_iso": "2026-09-27"}])["success"]
    with SessionLocal() as session:
        artifact = session.get(Artifact, md.id)
        written = dt.commit_version(session, artifact)
        previous = _previous_version(session, artifact, written.version)
        restored = dt.restore_version(session, artifact, previous)
        session.commit()
        versions = session.scalars(
            select(ArtifactVersion)
            .where(ArtifactVersion.artifact_id == artifact.id)
            .order_by(ArtifactVersion.version)
        ).all()
    assert restored["success"] is True
    assert [version.version for version in versions] == [1, 2, 3]


def test_ambiguous_suggestion_is_marked_unsafe(artifacts):
    md = artifacts["md"]
    outcome = dt.generate_patch(
        md,
        [{"location": "line_1", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}],
        [{"location": "line_1", "text": "两个日期 9月20日 和 9月20日", "meta": {}}],
    )
    assert outcome["success"] is True
    assert outcome["result"]["safe_to_apply"] is False
    assert outcome["result"]["ambiguous_edits"] == 1


def test_suggestion_patches_are_unique_per_action(artifacts):
    docx_art = artifacts["docx"]
    blocks = dt.parse_file(docx_art.source_path, "docx")["blocks"]
    edits = [{"location": "para_0", "old_iso": "2026-09-20", "new_iso": "2026-09-27"}]
    first = dt.generate_patch(docx_art, edits, blocks, patch_id="act_first")
    second = dt.generate_patch(docx_art, edits, blocks, patch_id="act_second")
    first_path = Path(first["result"]["patch_path"])
    second_path = Path(second["result"]["patch_path"])
    assert first_path != second_path
    assert first_path.exists() and second_path.exists()
