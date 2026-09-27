"""Parser tests: markdown, docx, xlsx with stable locations."""
from pathlib import Path

import pytest

from backend.parsers import parse_file
from backend.parsers.markdown_parser import parse_text


def test_markdown_blocks_and_locations():
    parsed = parse_text("# 标题\n\n第一行内容\n第二行 9 月 20 日\n")
    locations = [b["location"] for b in parsed["blocks"]]
    assert locations == ["line_1", "line_3", "line_4"]
    assert parsed["blocks"][1]["text"] == "第一行内容"


def test_xlsx_key_value_rows(tmp_path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(("Project", "Alpha V2.0"))
    ws.append(("Release Date", "2026-09-20"))
    path = tmp_path / "plan.xlsx"
    wb.save(path)

    parsed = parse_file(path)
    assert parsed["artifact_type"] == "xlsx"
    row2 = parsed["blocks"][1]
    assert row2["text"] == "Release Date: 2026-09-20"
    assert row2["meta"]["context_entity"] == "Alpha V2.0"
    assert row2["location"].startswith("Sheet!") or "Sheet" in row2["location"]


def test_xlsx_context_entity_not_picked_from_numeric_value(tmp_path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(("Release Date", "2026-09-20"))
    path = tmp_path / "plan.xlsx"
    wb.save(path)
    parsed = parse_file(path)
    assert parsed["blocks"][0]["meta"]["context_entity"] == ""


def test_docx_paragraphs_and_tables(tmp_path: Path):
    import docx

    document = docx.Document()
    document.add_paragraph("Alpha V2.0 将于 2026 年 9 月 20 日正式发布。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Release Date"
    table.rows[0].cells[1].text = "2026-09-20"
    path = tmp_path / "PRD.docx"
    document.save(path)

    parsed = parse_file(path)
    texts = [b["text"] for b in parsed["blocks"]]
    assert any("2026 年 9 月 20 日" in t for t in texts)
    assert any("Release Date | 2026-09-20" in t for t in texts)


def test_unsupported_type_raises(tmp_path: Path):
    with pytest.raises(ValueError):
        parse_file(tmp_path / "file.pdf", kind="pdf")
